from fastapi import FastAPI, HTTPException, Depends, UploadFile, File, Header
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, StreamingResponse
import psycopg2
import psycopg2.extras
import sqlite3
import hashlib, secrets, os, io, json
from datetime import datetime, timedelta
from typing import Optional, List
from calendar import monthrange, weekday as cal_weekday
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import imaplib
import email as email_lib
from email.header import decode_header as decode_email_header
import threading
import time

# ─── Email настройки (из Railway Variables) ───────────────────────────────────
EMAIL_HOST     = os.environ.get("EMAIL_HOST", "outlook.office365.com")
EMAIL_PORT     = int(os.environ.get("EMAIL_PORT", "993"))
EMAIL_USER     = os.environ.get("EMAIL_USER", "")
EMAIL_PASSWORD = os.environ.get("EMAIL_PASSWORD", "")
EMAIL_SENDER   = os.environ.get("EMAIL_SENDER", "reports@kari.com")
EMAIL_SUBJECT  = os.environ.get("EMAIL_SUBJECT", "Мотивация KariKids")
EMAIL_HOUR     = int(os.environ.get("EMAIL_FETCH_HOUR", "9"))
EMAIL_MINUTE   = int(os.environ.get("EMAIL_FETCH_MINUTE", "0"))

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

DATABASE_URL = os.environ.get("DATABASE_URL", "")
ADMIN_PIN = "5712"  # Фиксированный PIN администратора

# ─── Справочник магазинов ────────────────────────────────────────────────────
STORE_DATA = {
    15159: {"name": "Пушкино Парк",    "plan": 7149435, "sr": 60000, "ar": 70000,  "dr": 100000, "dc": 1, "ac": 2, "sc": 3, "nc": 0,   "cl": 30000, "ld": 6000},
    15003: {"name": "Чехов Карнавал",  "plan": 3989304, "sr": 55000, "ar": 65000,  "dr": 90000,  "dc": 1, "ac": 1, "sc": 2, "nc": 0.5, "cl": 15000, "ld": 7000},
    15023: {"name": "Раменское",       "plan": 3876988, "sr": 55000, "ar": 65000,  "dr": 95000,  "dc": 1, "ac": 2, "sc": 2, "nc": 0.5, "cl": 18000, "ld": 11250},
    15171: {"name": "Видное Галерея",  "plan": 3222628, "sr": 60000, "ar": 65000,  "dr": 90000,  "dc": 1, "ac": 1, "sc": 2, "nc": 0,   "cl": 27000, "ld": 9000},
    15141: {"name": "Жуковский",       "plan": 3017359, "sr": 60000, "ar": 60000,  "dr": 90000,  "dc": 1, "ac": 1, "sc": 2, "nc": 0,   "cl": 15000, "ld": 9000},
    15147: {"name": "Янтарь",          "plan": 2626992, "sr": 60000, "ar": 70000,  "dr": 90000,  "dc": 1, "ac": 1, "sc": 2, "nc": 0,   "cl": 27000, "ld": 9000},
    15222: {"name": "ТЦ Круг",         "plan": 2596560, "sr": 60000, "ar": 70000,  "dr": 95000,  "dc": 1, "ac": 1, "sc": 2, "nc": 0,   "cl": 30000, "ld": 9000},
    15145: {"name": "Серпухов Атлас",  "plan": 2276176, "sr": 50000, "ar": 60000,  "dr": 90000,  "dc": 1, "ac": 1, "sc": 2, "nc": 0,   "cl": 25000, "ld": 9000},
    15170: {"name": "Коломна КАДО",    "plan": 2214011, "sr": 50000, "ar": 65000,  "dr": 85000,  "dc": 1, "ac": 1, "sc": 2, "nc": 0,   "cl": 30000, "ld": 9000},
    15191: {"name": "ТЦ Облака",       "plan": 2004742, "sr": 60000, "ar": 70000,  "dr": 95000,  "dc": 1, "ac": 1, "sc": 2, "nc": 0,   "cl": 29000, "ld": 9000},
    15203: {"name": "Ивантеевка Твид", "plan": 1542018, "sr": 60000, "ar": 75000,  "dr": 90000,  "dc": 1, "ac": 0, "sc": 2, "nc": 0,   "cl": 0,     "ld": 0},
    15221: {"name": "Кузьминки Молл",  "plan": 1512887, "sr": 60000, "ar": 75000,  "dr": 90000,  "dc": 1, "ac": 1, "sc": 2, "nc": 0,   "cl": 29000, "ld": 9000},
}

# ─── Вспомогательные функции для данных магазина ─────────────────────────────
def get_store_ref_from_conn(store_id: int, conn) -> dict:
    """Возвращает плановые данные из fot_plan (последний месяц) или STORE_DATA"""
    row = conn.execute(
        "SELECT * FROM fot_plan WHERE store_id=? ORDER BY month DESC LIMIT 1",
        (store_id,)
    ).fetchone()
    if row:
        base = STORE_DATA.get(store_id, {})
        return {
            "name": base.get("name", f"Магазин {store_id}"),
            "plan": int(row["plan"] or 0),
            "sr":   int(row["sr"]   or 0),
            "ar":   int(row["ar"]   or 0),
            "dr":   int(row["dr"]   or 0),
            "dc":   float(row["dc"] or 0),
            "ac":   float(row["ac"] or 0),
            "sc":   float(row["sc"] or 0),
            "nc":   float(row["nc"] or 0),
            "cl":       int(row["cl"]        or 0),
            "ld":       int(row["ld"]        or 0),
            "fot_total": int(row["fot_total"] or 0) if "fot_total" in row else 0,
        }
    return STORE_DATA.get(store_id)

# ─── БД (PostgreSQL) ─────────────────────────────────────────────────────────
class DBConn:
    """Wrapper supporting both SQLite (no DATABASE_URL) and PostgreSQL."""
    def __init__(self, dsn):
        if dsn:
            self._conn = psycopg2.connect(dsn)
            self._is_pg = True
        else:
            self._conn = sqlite3.connect("fot.db", check_same_thread=False)
            self._conn.row_factory = lambda cur, row: {
                col[0]: row[idx] for idx, col in enumerate(cur.description)
            } if cur.description else {}
            self._is_pg = False

    def execute(self, sql, params=None):
        if self._is_pg:
            sql = sql.replace("?", "%s")
            sql = sql.replace("INTEGER PRIMARY KEY", "SERIAL PRIMARY KEY")
            cur = self._conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        else:
            cur = self._conn.cursor()
        cur.execute(sql, params) if params else cur.execute(sql)
        return cur

    def commit(self):
        self._conn.commit()

    def rollback(self):
        self._conn.rollback()

    def close(self):
        self._conn.close()

def get_db() -> DBConn:
    return DBConn(DATABASE_URL)

def init_db():
    conn = get_db()
    conn.execute("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY,
            name TEXT NOT NULL,
            pin_hash TEXT NOT NULL,
            store_id INTEGER,
            role TEXT DEFAULT 'director',
            token TEXT,
            token_expires TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS user_sessions (
            id INTEGER PRIMARY KEY,
            user_id INTEGER NOT NULL,
            token TEXT NOT NULL UNIQUE,
            token_expires TEXT NOT NULL,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS motivation_data (
            id INTEGER PRIMARY KEY,
            store_id INTEGER NOT NULL,
            report_date TEXT NOT NULL,
            login TEXT,
            role TEXT,
            name TEXT,
            to_fact REAL DEFAULT 0,
            pct_to REAL DEFAULT 0,
            income REAL DEFAULT 0,
            fdm REAL DEFAULT 0,
            is_total INTEGER DEFAULT 0,
            is_bezshk INTEGER DEFAULT 0,
            uploaded_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS upload_log (
            id INTEGER PRIMARY KEY,
            filename TEXT,
            report_date TEXT,
            rows_count INTEGER,
            uploaded_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS fot_plan (
            id INTEGER PRIMARY KEY,
            store_id INTEGER NOT NULL,
            month TEXT NOT NULL,
            plan INTEGER DEFAULT 0,
            sr INTEGER DEFAULT 0,
            ar INTEGER DEFAULT 0,
            dr INTEGER DEFAULT 0,
            dc REAL DEFAULT 0,
            ac REAL DEFAULT 0,
            sc REAL DEFAULT 0,
            nc REAL DEFAULT 0,
            cl INTEGER DEFAULT 0,
            ld INTEGER DEFAULT 0,
            uploaded_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS fot_upload_log (
            id INTEGER PRIMARY KEY,
            filename TEXT,
            month TEXT,
            stores_count INTEGER,
            uploaded_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS work_schedule (
            id INTEGER PRIMARY KEY,
            store_id INTEGER NOT NULL,
            month TEXT NOT NULL,
            employee_name TEXT NOT NULL,
            employee_role TEXT DEFAULT '',
            days TEXT DEFAULT '[]',
            updated_at TEXT DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(store_id, month, employee_name)
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS sales_data (
            id INTEGER PRIMARY KEY,
            report_date TEXT NOT NULL,
            store_id INTEGER,
            subdivision TEXT,
            to_fact REAL DEFAULT 0,
            plan_pct REAL DEFAULT 0,
            lfl REAL DEFAULT 0,
            margin REAL DEFAULT 0,
            avg_ticket REAL DEFAULT 0,
            conversion REAL DEFAULT 0,
            upt REAL DEFAULT 0,
            traffic INTEGER DEFAULT 0,
            traffic_lfl REAL DEFAULT 0,
            toys_to REAL DEFAULT 0,    toys_lfl REAL DEFAULT 0,
            clothes_to REAL DEFAULT 0, clothes_lfl REAL DEFAULT 0,
            shoes_to REAL DEFAULT 0,   shoes_lfl REAL DEFAULT 0,
            sport_to REAL DEFAULT 0,   sport_lfl REAL DEFAULT 0,
            uploaded_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS sales_upload_log (
            id INTEGER PRIMARY KEY,
            filename TEXT,
            report_date TEXT,
            rows_count INTEGER,
            uploaded_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS stock_data (
            id INTEGER PRIMARY KEY,
            report_date TEXT NOT NULL,
            store_id INTEGER,
            subdivision TEXT,
            total_stock REAL DEFAULT 0,
            total_lfl REAL DEFAULT 0,
            clothes_stock REAL DEFAULT 0,
            clothes_lfl REAL DEFAULT 0,
            toys_stock REAL DEFAULT 0,
            toys_lfl REAL DEFAULT 0,
            shoes_stock REAL DEFAULT 0,
            shoes_lfl REAL DEFAULT 0,
            sport_stock REAL DEFAULT 0,
            sport_lfl REAL DEFAULT 0,
            uploaded_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS stock_upload_log (
            id INTEGER PRIMARY KEY,
            filename TEXT,
            report_date TEXT,
            rows_count INTEGER,
            uploaded_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.commit()
    # Добавляем колонку first_login если нет (миграция)
    try:
        conn.execute("ALTER TABLE users ADD COLUMN first_login INTEGER DEFAULT 1")
        conn.commit()
    except Exception:
        conn.rollback()
    try:
        conn.execute("ALTER TABLE fot_plan ADD COLUMN fot_total INTEGER DEFAULT 0")
        conn.commit()
    except Exception:
        conn.rollback()  # Колонка уже существует

    # Создаём или обновляем аккаунт администратора
    pin_hash = hashlib.sha256(ADMIN_PIN.encode()).hexdigest()
    if not conn.execute("SELECT id FROM users WHERE role='admin'").fetchone():
        conn.execute(
            "INSERT INTO users (name, pin_hash, store_id, role, first_login) VALUES (?,?,NULL,'admin',0)",
            ("Администратор", pin_hash)
        )
    else:
        # Принудительно обновляем хэш при каждом старте (подхватывает новый ADMIN_PIN)
        conn.execute("UPDATE users SET pin_hash=? WHERE role='admin'", (pin_hash,))
    conn.commit()

    # Авто-создаём аккаунты директоров для всех магазинов
    default_pin_hash = hashlib.sha256("1111".encode()).hexdigest()
    for sid, sdata in STORE_DATA.items():
        if not conn.execute("SELECT id FROM users WHERE store_id=? AND role='director'", (sid,)).fetchone():
            conn.execute(
                "INSERT INTO users (name, pin_hash, store_id, role, first_login) VALUES (?,?,?,'director',1)",
                (f"Магазин {sid}", default_pin_hash, sid)
            )
    conn.commit()
    conn.close()

init_db()

# ─── Утилиты ─────────────────────────────────────────────────────────────────
def hash_pin(pin: str) -> str:
    return hashlib.sha256(pin.strip().encode()).hexdigest()

def plan_coef(pct: float) -> float:
    coef = pct / 100.0
    return min(round(coef, 4), 1.10)

def get_user_by_token(token: str):
    if not token:
        return None
    try:
        conn = get_db()
        # Ищем в таблице сессий (новый механизм)
        row = conn.execute("""
            SELECT u.* FROM users u
            JOIN user_sessions s ON s.user_id = u.id
            WHERE s.token=? AND s.token_expires > ?
        """, (token, datetime.now().isoformat())).fetchone()
        if not row:
            # Fallback: старый механизм (один токен в users)
            row = conn.execute(
                "SELECT * FROM users WHERE token=? AND token_expires > ?",
                (token, datetime.now().isoformat())
            ).fetchone()
        conn.close()
        return row
    except Exception:
        return None

def require_auth(authorization: Optional[str] = Header(None)):
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Не авторизован")
    token = authorization.split(" ")[1]
    user = get_user_by_token(token)
    if not user:
        raise HTTPException(status_code=401, detail="Сессия истекла")
    return dict(user)

def require_admin(user=Depends(require_auth)):
    if user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Нет доступа")
    return user

# ─── Auth ────────────────────────────────────────────────────────────────────
@app.post("/api/login")
def login(body: dict):
    pin = body.get("pin", "").strip()
    store_id_raw = body.get("store_id")
    if not pin:
        raise HTTPException(status_code=400, detail="Введите пин-код")
    pin_hash = hash_pin(pin)
    conn = get_db()
    if store_id_raw:
        # Вход директора: по store_id + PIN
        try:
            sid = int(store_id_raw)
        except (ValueError, TypeError):
            conn.close()
            raise HTTPException(status_code=400, detail="Неверный номер магазина")
        user = conn.execute(
            "SELECT * FROM users WHERE store_id=? AND role='director' AND pin_hash=?",
            (sid, pin_hash)
        ).fetchone()
        if not user:
            conn.close()
            raise HTTPException(status_code=401, detail="Неверный номер магазина или пин-код")
    else:
        # Вход администратора: только PIN
        user = conn.execute(
            "SELECT * FROM users WHERE role='admin' AND pin_hash=?",
            (pin_hash,)
        ).fetchone()
        if not user:
            conn.close()
            raise HTTPException(status_code=401, detail="Неверный пин-код")
    token = secrets.token_hex(32)
    expires = (datetime.now() + timedelta(days=90)).isoformat()
    conn.execute(
        "INSERT INTO user_sessions (user_id, token, token_expires) VALUES (?,?,?)",
        (user["id"], token, expires)
    )
    conn.execute("UPDATE users SET token=?, token_expires=? WHERE id=?",
                 (token, expires, user["id"]))
    conn.commit()
    first_login = bool(user.get("first_login", 0))
    conn.close()
    return {
        "token": token,
        "name": user["name"],
        "role": user["role"],
        "store_id": user["store_id"],
        "first_login": first_login,
        "store_name": STORE_DATA.get(user["store_id"], {}).get("name") if user["store_id"] else None
    }

@app.post("/api/change_pin")
def change_pin(body: dict, user=Depends(require_auth)):
    new_pin = body.get("new_pin", "").strip()
    if not new_pin or len(new_pin) < 4:
        raise HTTPException(status_code=400, detail="Пин-код минимум 4 цифры")
    if not new_pin.isdigit():
        raise HTTPException(status_code=400, detail="Пин-код должен состоять только из цифр")
    new_hash = hash_pin(new_pin)
    conn = get_db()
    conn.execute("UPDATE users SET pin_hash=?, first_login=0 WHERE id=?", (new_hash, user["id"]))
    conn.commit()
    conn.close()
    return {"ok": True}

@app.post("/api/reset_pin/{user_id}")
def reset_pin(user_id: int, admin=Depends(require_admin)):
    default_hash = hashlib.sha256("1111".encode()).hexdigest()
    conn = get_db()
    row = conn.execute("SELECT id FROM users WHERE id=? AND role='director'", (user_id,)).fetchone()
    if not row:
        conn.close()
        raise HTTPException(status_code=404, detail="Директор не найден")
    conn.execute("UPDATE users SET pin_hash=?, first_login=1 WHERE id=?", (default_hash, user_id))
    conn.commit()
    conn.close()
    return {"ok": True}

@app.get("/api/system/fix-admin-5712-now")
def fix_admin_pin():
    """Временный endpoint: принудительно сбрасывает PIN админа на 5712"""
    ph = hashlib.sha256("5712".encode()).hexdigest()
    conn = get_db()
    conn.execute("UPDATE users SET pin_hash=? WHERE role='admin'", (ph,))
    conn.commit()
    rows = conn.execute("SELECT id, name, role FROM users WHERE role='admin'").fetchall()
    conn.close()
    return {"ok": True, "admins_updated": [dict(r) for r in rows]}

@app.post("/api/logout")
def logout(user=Depends(require_auth)):
    conn = get_db()
    conn.execute("UPDATE users SET token=NULL WHERE id=?", (user["id"],))
    conn.commit()
    conn.close()
    return {"ok": True}

@app.get("/api/me")
def me(user=Depends(require_auth)):
    return {
        "name": user["name"],
        "role": user["role"],
        "store_id": user["store_id"],
        "store_name": STORE_DATA.get(user["store_id"], {}).get("name") if user["store_id"] else None
    }

# ─── Загрузка файла мотивации (только admin) ─────────────────────────────────
@app.post("/api/upload")
async def upload_motivation(file: UploadFile = File(...), user=Depends(require_admin)):
    content = await file.read()
    try:
        df = pd.read_excel(io.BytesIO(content), header=None)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Ошибка чтения файла: {e}")

    header_row = None
    for i, row in df.iterrows():
        if str(row[0]).strip().startswith("Дата"):
            header_row = i
            break
    if header_row is None:
        raise HTTPException(status_code=400, detail="Не найден заголовок 'Дата отчета'")

    data_df = df.iloc[header_row + 1:].reset_index(drop=True)
    conn = get_db()
    rows_inserted = 0
    report_date = None
    cleared_stores = set()  # stores whose old data has been deleted

    for _, row in data_df.iterrows():
        try:
            store_id = int(float(str(row[2]))) if str(row[2]) not in ["nan", ""] else None
            if not store_id:
                continue
            login = str(row[3] or "").strip()
            role = str(row[5] or "").strip()
            name = str(row[6] or "").strip()
            to_fact = float(row[7]) if str(row[7]) not in ["nan", ""] else 0
            pct_to = float(row[8]) if str(row[8]) not in ["nan", ""] else 0
            income = float(row[9]) if str(row[9]) not in ["nan", ""] else 0
            fdm_val = float(row[10]) if str(row[10]) not in ["nan", ""] else 0

            raw_date = row[0]
            if hasattr(raw_date, "strftime"):
                rd = raw_date.strftime("%d.%m.%Y")
            elif isinstance(raw_date, (int, float)):
                from xlrd import xldate_as_datetime
                rd = xldate_as_datetime(raw_date, 0).strftime("%d.%m.%Y")
            else:
                rd = str(raw_date).split(" ")[0].strip()
            if report_date is None:
                report_date = rd

            # При первом появлении магазина — удаляем его старые данные (обновление без дублей)
            if store_id not in cleared_stores:
                conn.execute("DELETE FROM motivation_data WHERE store_id=?", (store_id,))
                cleared_stores.add(store_id)

            is_total = 1 if login == "Total" else 0
            is_bezshk = 1 if login == "БезШК" else 0

            conn.execute("""
                INSERT INTO motivation_data
                (store_id, report_date, login, role, name, to_fact, pct_to, income, fdm, is_total, is_bezshk)
                VALUES (?,?,?,?,?,?,?,?,?,?,?)
            """, (store_id, rd, login, role, name, to_fact, pct_to, income, fdm_val, is_total, is_bezshk))
            rows_inserted += 1
        except Exception:
            continue

    conn.execute("INSERT INTO upload_log (filename, report_date, rows_count) VALUES (?,?,?)",
                 (file.filename, report_date, rows_inserted))
    conn.commit()
    conn.close()
    return {"ok": True, "rows": rows_inserted, "date": report_date}

# ─── Данные магазина ──────────────────────────────────────────────────────────
@app.get("/api/store/{store_id}")
def get_store_data(store_id: int, days_total: int = 31, forecast_pct: int = 100, mode: str = "avg", user=Depends(require_auth)):
    if user["role"] == "director" and user["store_id"] != store_id:
        raise HTTPException(status_code=403, detail="Нет доступа к этому магазину")

    conn = get_db()
    ref = get_store_ref_from_conn(store_id, conn)
    if not ref:
        conn.close()
        raise HTTPException(status_code=404, detail="Магазин не найден")
    latest = conn.execute(
        "SELECT MAX(report_date) as d FROM motivation_data WHERE store_id=?", (store_id,)
    ).fetchone()
    if not latest or not latest["d"]:
        conn.close()
        return {"store_id": store_id, "name": ref["name"], "no_data": True}

    report_date = latest["d"]
    day_report = int(report_date.split(".")[0])

    rows = conn.execute(
        "SELECT * FROM motivation_data WHERE store_id=? AND report_date=?",
        (store_id, report_date)
    ).fetchall()

    # Загружаем часы из графика работы за текущий месяц
    STANDARD_HOURS = 168
    parts = report_date.split(".")
    month_str = f"{parts[2]}-{parts[1]}"  # YYYY-MM
    sched_rows = conn.execute(
        "SELECT employee_name, days FROM work_schedule WHERE store_id=? AND month=?",
        (store_id, month_str)
    ).fetchall()
    conn.close()

    hours_map = {}
    for srow in sched_rows:
        try:
            days_list = json.loads(srow["days"] or "[]")
            # Считаем заполненным только если есть хоть одна непустая ячейка
            if any(d not in ("", None) for d in days_list):
                total_h = sum(int(d) for d in days_list if str(d).isdigit() and int(d) > 0)
                hours_map[srow["employee_name"]] = total_h  # сохраняем даже если 0
        except Exception:
            pass

    def get_hours_coef(name):
        if name not in hours_map:
            return 1.0  # нет данных графика → стандартный коэф
        h = hours_map[name]
        return round(h / STANDARD_HOURS, 4) if h > 0 else 0.0  # заполнен, но 0 часов → коэф 0

    total_row = next((r for r in rows if r["is_total"]), None)
    staff = [r for r in rows if not r["is_total"] and not r["is_bezshk"] and r["role"] and r["role"] != "НетДолжности"]

    fact_to = total_row["to_fact"] if total_row else 0
    fact_fot = (total_row["income"] + total_row["fdm"]) if total_row else 0

    ratio = days_total / day_report
    forecast_to = fact_to * ratio if mode == "avg" else ref["plan"] * forecast_pct / 100
    plan_pct = (forecast_to / ref["plan"] * 100) if ref["plan"] > 0 else 0
    fact_plan_pct = (fact_to / ref["plan"] * 100) if ref["plan"] > 0 else 0
    coef = plan_coef(plan_pct)

    extra = ref["cl"] + ref["ld"]  # клининг + погрузка (только для справки)

    def get_rate(role):
        if "Директор" in role: return ref["dr"]
        if "Администратор" in role: return ref["ar"]
        return ref["sr"]

    staff_list = sorted([{
        "name": r["name"],
        "role": r["role"],
        "rate": get_rate(r["role"]),
        "fact_income": r["income"],
        "hours": hours_map.get(r["name"], 0),
        "hours_coef": get_hours_coef(r["name"]),
        "forecast_salary": round(get_rate(r["role"]) * coef * get_hours_coef(r["name"])),
    } for r in staff], key=lambda x: (
        0 if "Директор" in x["role"] else 1 if "Администратор" in x["role"] else 2
    ))

    # Плановый ФОТ = "ЗП ИТОГО" из файла ФОТ
    planned_fot = ref.get("fot_total", 0) or 0
    # Если данных из файла нет — считаем по ставкам как запасной вариант
    if planned_fot <= 0:
        planned_fot = sum(get_rate(r["role"]) for r in staff)

    # Максимально допустимый ФОТ = плановый × коэф плана (макс 110%)
    max_fot = round(planned_fot * coef)

    # Расчётные зарплаты сотрудников
    fot_wages_raw = sum(s["forecast_salary"] for s in staff_list)

    # Если сумма превышает лимит — пропорционально уменьшаем каждому
    if fot_wages_raw > max_fot and fot_wages_raw > 0:
        scale = max_fot / fot_wages_raw
        for s in staff_list:
            s["forecast_salary"] = round(s["forecast_salary"] * scale)

    fot_wages = sum(s["forecast_salary"] for s in staff_list)
    forecast_fot = fot_wages
    fot_capped = fot_wages_raw > max_fot

    rest_fot = max(0, forecast_fot - fact_fot)

    return {
        "store_id": store_id,
        "name": ref["name"],
        "report_date": report_date,
        "plan": ref["plan"],
        "fact_to": fact_to,
        "fact_fot": fact_fot,
        "fact_plan_pct": round(fact_plan_pct, 1),
        "forecast_to": round(forecast_to),
        "forecast_fot": round(forecast_fot),
        "forecast_plan_pct": round(plan_pct, 1),
        "coef": coef,
        "fot_wages": round(fot_wages),
        "planned_fot": round(planned_fot),
        "budget_fot": round(planned_fot),  # backward compat
        "max_fot": round(max_fot),
        "fot_capped": fot_capped,
        "extra": extra,
        "rest_fot": round(rest_fot),
        "fot_to_pct": round(forecast_fot / forecast_to * 100, 1) if forecast_to > 0 else 0,
        "staff": staff_list,
    }

# ─── Сводка по всем магазинам (только admin) ──────────────────────────────────
@app.get("/api/summary")
def get_summary(days_total: int = 31, forecast_pct: int = 100, mode: str = "avg", user=Depends(require_auth)):
    if user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Нет доступа")
    result = []
    for sid in STORE_DATA:
        try:
            data = get_store_data(sid, days_total, forecast_pct, mode, user)
            if not data.get("no_data"):
                result.append(data)
        except Exception:
            continue
    result.sort(key=lambda x: x.get("forecast_to", 0), reverse=True)
    return result

@app.get("/api/admin/uploads")
def get_summary_compat(days_total: int = 31, forecast_pct: int = 100, mode: str = "avg", user=Depends(require_auth)):
    """Compatibility alias for old frontend"""
    if user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Нет доступа")
    result = []
    for sid in STORE_DATA:
        try:
            data = get_store_data(sid, days_total, forecast_pct, mode, user)
            if not data.get("no_data"):
                data["store_name"] = data.get("name", "")
                result.append(data)
        except Exception:
            continue
    result.sort(key=lambda x: x.get("forecast_to", 0), reverse=True)
    return result

# ─── Управление пользователями (только admin) ─────────────────────────────────
@app.get("/api/users")
def get_users(user=Depends(require_admin)):
    conn = get_db()
    users = conn.execute("SELECT id, name, store_id, role, first_login, created_at FROM users").fetchall()
    conn.close()
    return [{"id": u["id"], "name": u["name"], "store_id": u["store_id"],
             "role": u["role"], "first_login": bool(u.get("first_login", 0)),
             "store_name": STORE_DATA.get(u["store_id"], {}).get("name") if u["store_id"] else None,
             "created_at": str(u["created_at"])} for u in users]

@app.post("/api/users")
def create_user(body: dict, user=Depends(require_admin)):
    name = body.get("name", "").strip()
    pin = body.get("pin", "").strip()
    store_id = body.get("store_id")
    role = body.get("role", "director")
    if not name or not pin:
        raise HTTPException(status_code=400, detail="Имя и пин-код обязательны")
    if len(pin) < 4:
        raise HTTPException(status_code=400, detail="Пин-код минимум 4 символа")
    pin_hash = hash_pin(pin)
    conn = get_db()
    existing = conn.execute("SELECT id FROM users WHERE pin_hash=?", (pin_hash,)).fetchone()
    if existing:
        conn.close()
        raise HTTPException(status_code=400, detail="Такой пин-код уже существует")
    conn.execute("INSERT INTO users (name, pin_hash, store_id, role) VALUES (?,?,?,?)",
                 (name, pin_hash, store_id, role))
    conn.commit()
    conn.close()
    return {"ok": True}

@app.delete("/api/users/{user_id}")
def delete_user(user_id: int, user=Depends(require_admin)):
    conn = get_db()
    target = conn.execute("SELECT role FROM users WHERE id=?", (user_id,)).fetchone()
    if not target:
        conn.close()
        raise HTTPException(status_code=404, detail="Пользователь не найден")
    if target["role"] == "admin":
        conn.close()
        raise HTTPException(status_code=400, detail="Нельзя удалить администратора")
    conn.execute("DELETE FROM users WHERE id=?", (user_id,))
    conn.commit()
    conn.close()
    return {"ok": True}

@app.put("/api/users/{user_id}/pin")
def change_pin(user_id: int, body: dict, user=Depends(require_admin)):
    new_pin = body.get("pin", "").strip()
    if len(new_pin) < 4:
        raise HTTPException(status_code=400, detail="Пин-код минимум 4 символа")
    conn = get_db()
    conn.execute("UPDATE users SET pin_hash=? WHERE id=?", (hash_pin(new_pin), user_id))
    conn.commit()
    conn.close()
    return {"ok": True}

@app.get("/api/upload_log")
def upload_log(user=Depends(require_admin)):
    conn = get_db()
    logs = conn.execute("SELECT * FROM upload_log ORDER BY id DESC LIMIT 10").fetchall()
    conn.close()
    return [dict(l) for l in logs]

# ─── Загрузка файла ФОТ (месячный план) ──────────────────────────────────────
@app.post("/api/fot_upload")
async def upload_fot(file: UploadFile = File(...), month: str = "", user=Depends(require_admin)):
    if not month:
        raise HTTPException(status_code=400, detail="Укажите месяц (YYYY-MM)")
    content = await file.read()
    try:
        df = pd.read_excel(io.BytesIO(content), header=None)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Ошибка чтения файла: {e}")

    # Находим столбец со store_id (5-значный числовой ID магазина)
    known_ids = set(STORE_DATA.keys())
    sid_col = None
    for ci in range(min(6, len(df.columns))):
        for val in df.iloc[:, ci]:
            try:
                v = int(float(str(val)))
                if v in known_ids:
                    sid_col = ci
                    break
            except Exception:
                pass
        if sid_col is not None:
            break
    if sid_col is None:
        raise HTTPException(status_code=400, detail="Не найдены ID магазинов. Проверьте формат файла.")

    def safe_int(val):
        try: return int(float(str(val)))
        except: return 0

    def safe_float(val):
        try: return float(str(val))
        except: return 0.0

    conn = get_db()
    stores_inserted = 0

    for _, row in df.iterrows():
        try:
            sid = int(float(str(row.iloc[sid_col])))
            if sid not in known_ids:
                continue
        except Exception:
            continue

        # Удаляем старую запись за этот месяц для данного магазина
        conn.execute("DELETE FROM fot_plan WHERE store_id=? AND month=?", (sid, month))

        conn.execute("""
            INSERT INTO fot_plan
            (store_id, month, dc, ac, sc, nc, plan, sr, ar, dr, fot_total, cl, ld)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
        """, (
            sid, month,
            safe_float(row.iloc[sid_col + 4]),   # Директор (кол-во)
            safe_float(row.iloc[sid_col + 5]),   # Администратор (кол-во)
            safe_float(row.iloc[sid_col + 6]),   # Продавец штат (кол-во)
            safe_float(row.iloc[sid_col + 7]),   # Продавец наймикс (кол-во)
            safe_int(row.iloc[sid_col + 8]),     # бюджет ТО (план)
            safe_int(row.iloc[sid_col + 9]),     # Средняя ЗП Продавца
            safe_int(row.iloc[sid_col + 10]),    # Зарплата АМ
            safe_int(row.iloc[sid_col + 11]),    # Зарплата ДМ
            safe_int(row.iloc[sid_col + 12]),    # ЗП ИТОГО (плановый ФОТ)
            safe_int(row.iloc[sid_col + 18]),    # Клининг
            safe_int(row.iloc[sid_col + 20]),    # Разгрузка погрузка
        ))
        stores_inserted += 1

    conn.execute(
        "INSERT INTO fot_upload_log (filename, month, stores_count) VALUES (?,?,?)",
        (file.filename, month, stores_inserted)
    )
    conn.commit()
    conn.close()
    return {"ok": True, "stores_count": stores_inserted, "month": month}

@app.get("/api/fot_log")
def fot_log(user=Depends(require_admin)):
    conn = get_db()
    logs = conn.execute("SELECT * FROM fot_upload_log ORDER BY id DESC LIMIT 10").fetchall()
    conn.close()
    return [dict(l) for l in logs]

@app.get("/api/fot_active_month")
def fot_active_month(user=Depends(require_admin)):
    conn = get_db()
    row = conn.execute("SELECT month FROM fot_plan ORDER BY month DESC LIMIT 1").fetchone()
    conn.close()
    return {"month": row["month"] if row else None}

@app.get("/api/stores")
def get_stores_list(user=Depends(require_admin)):
    return [{"id": k, "name": v["name"]} for k, v in STORE_DATA.items()]

@app.get("/api/debug/store_ids")
def debug_store_ids():
    conn = get_db()
    rows = conn.execute(
        "SELECT DISTINCT store_id, report_date, COUNT(*) as cnt FROM motivation_data GROUP BY store_id, report_date"
    ).fetchall()
    conn.close()
    known = list(STORE_DATA.keys())
    return {
        "in_db": [{"store_id": r["store_id"], "date": r["report_date"], "rows": r["cnt"], "in_store_data": r["store_id"] in known} for r in rows],
        "store_data_ids": known
    }

# ─── График работы ────────────────────────────────────────────────────────────
@app.get("/api/schedule/{store_id}/{month}")
def get_schedule(store_id: int, month: str, user=Depends(require_auth)):
    if user["role"] == "director" and user["store_id"] != store_id:
        raise HTTPException(status_code=403, detail="Нет доступа")
    try:
        year, mon = map(int, month.split("-"))
        _, days_in_month = monthrange(year, mon)
    except Exception:
        raise HTTPException(status_code=400, detail="Неверный формат месяца (YYYY-MM)")

    try:
        conn = get_db()
        rows = conn.execute(
            "SELECT * FROM work_schedule WHERE store_id=? AND month=? ORDER BY id",
            (store_id, month)
        ).fetchall()
        # PostgreSQL: DISTINCT + ORDER BY CASE требует подзапрос
        employees_db = conn.execute("""
            SELECT name, role FROM (
                SELECT DISTINCT name, role
                FROM motivation_data
                WHERE store_id=?
                  AND is_total=0 AND is_bezshk=0
                  AND name IS NOT NULL AND name NOT IN ('', 'nan')
                  AND role IS NOT NULL AND role NOT IN ('', 'НетДолжности', 'nan')
            ) sub
            ORDER BY
                CASE
                    WHEN role LIKE '%%Директор%%' THEN 1
                    WHEN role LIKE '%%Администратор%%' THEN 2
                    ELSE 3
                END, name
        """, (store_id,)).fetchall()
        conn.close()
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"DB error: {str(e)}")

    schedule_map = {r["employee_name"]: json.loads(r["days"] or "[]") for r in rows}
    result = []
    for emp in employees_db:
        saved = schedule_map.get(emp["name"])
        if saved and len(saved) >= days_in_month:
            days = saved[:days_in_month]
        else:
            days = [""] * days_in_month
        result.append({"name": emp["name"], "role": emp["role"], "days": days})

    weekdays = [cal_weekday(year, mon, d + 1) for d in range(days_in_month)]

    return {
        "store_id": store_id,
        "month": month,
        "days_in_month": days_in_month,
        "weekdays": weekdays,
        "employees": result
    }


@app.post("/api/schedule/{store_id}/{month}")
def save_schedule(store_id: int, month: str, body: dict, user=Depends(require_auth)):
    if user["role"] == "director" and user["store_id"] != store_id:
        raise HTTPException(status_code=403, detail="Нет доступа")
    employees = body.get("employees", [])
    conn = get_db()
    saved = 0
    for emp in employees:
        name = str(emp.get("name", "")).strip()
        role = str(emp.get("role", ""))
        days = json.dumps(emp.get("days", []))
        if not name:
            continue
        conn.execute("""
            INSERT INTO work_schedule (store_id, month, employee_name, employee_role, days, updated_at)
            VALUES (%s,%s,%s,%s,%s,%s)
            ON CONFLICT (store_id, month, employee_name)
            DO UPDATE SET employee_role=EXCLUDED.employee_role,
                          days=EXCLUDED.days,
                          updated_at=EXCLUDED.updated_at
        """, (store_id, month, name, role, days, datetime.now().isoformat()))
        saved += 1
    conn.commit()
    conn.close()
    return {"ok": True, "saved": saved}


@app.get("/api/schedule/{store_id}/{month}/export")
def export_schedule(store_id: int, month: str, user=Depends(require_auth)):
    if user["role"] == "director" and user["store_id"] != store_id:
        raise HTTPException(status_code=403, detail="Нет доступа")

    data = get_schedule(store_id, month, user)
    store_name = STORE_DATA.get(store_id, {}).get("name", str(store_id))
    year, mon = map(int, month.split("-"))
    days_in_month = data["days_in_month"]
    weekdays = data["weekdays"]

    month_name = ["Январь","Февраль","Март","Апрель","Май","Июнь",
                  "Июль","Август","Сентябрь","Октябрь","Ноябрь","Декабрь"][mon - 1]

    wb = Workbook()
    ws = wb.active
    ws.title = "График"

    # Fills
    fill_weekend = PatternFill("solid", fgColor="D9D9D9")
    fill_work    = PatternFill("solid", fgColor="C6EFCE")
    fill_vac     = PatternFill("solid", fgColor="FFEB9C")
    fill_sick    = PatternFill("solid", fgColor="FFC7CE")
    fill_header  = PatternFill("solid", fgColor="1A56C0")
    fill_subhdr  = PatternFill("solid", fgColor="BDD7EE")

    bold_white = Font(bold=True, color="FFFFFF")
    bold_dark  = Font(bold=True)
    thin = Side(style="thin", color="AAAAAA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    # Title
    ws.merge_cells(f"A1:{get_column_letter(days_in_month + 4)}1")
    title_cell = ws["A1"]
    title_cell.value = f"График работы — {store_name} — {month_name} {year}"
    title_cell.font = bold_white
    title_cell.fill = fill_header
    title_cell.alignment = center
    ws.row_dimensions[1].height = 22

    # Column headers
    ws["A2"] = "№"
    ws["B2"] = "Сотрудник"
    ws["C2"] = "Должность"
    for d in range(1, days_in_month + 1):
        col = get_column_letter(d + 3)
        ws[f"{col}2"] = d
    ws[f"{get_column_letter(days_in_month + 4)}2"] = "Итого (ч)"

    for col in range(1, days_in_month + 5):
        cell = ws.cell(row=2, column=col)
        cell.fill = fill_subhdr
        cell.font = bold_dark
        cell.alignment = center
        cell.border = border
        # Weekend highlight in header
        if col > 3 and col < days_in_month + 4:
            if weekdays[col - 4] >= 5:
                cell.fill = fill_weekend

    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 18
    for d in range(1, days_in_month + 1):
        ws.column_dimensions[get_column_letter(d + 3)].width = 3.5
    ws.column_dimensions[get_column_letter(days_in_month + 4)].width = 6

    # Data rows
    fill_noshow = PatternFill("solid", fgColor="808080")
    STATUS_COLOR = {"О": fill_vac, "Б": fill_sick, "Н": fill_noshow}
    for i, emp in enumerate(data["employees"]):
        row = i + 3
        ws.cell(row=row, column=1, value=i + 1).border = border
        ws.cell(row=row, column=2, value=emp["name"]).border = border
        ws.cell(row=row, column=3, value=emp["role"]).border = border
        total_hours = 0
        for d in range(days_in_month):
            col = d + 4
            status = emp["days"][d] if d < len(emp["days"]) else ""
            # Числовые часы
            try:
                hours = int(status)
                cell = ws.cell(row=row, column=col, value=hours)
                cell.fill = fill_work
                total_hours += hours
            except (ValueError, TypeError):
                cell = ws.cell(row=row, column=col, value=status)
                if status == "В" or (not status and weekdays[d] >= 5):
                    cell.fill = fill_weekend
                elif status in STATUS_COLOR:
                    cell.fill = STATUS_COLOR[status]
            cell.alignment = center
            cell.border = border
        tot_cell = ws.cell(row=row, column=days_in_month + 4, value=total_hours)
        tot_cell.font = bold_dark
        tot_cell.alignment = center
        tot_cell.border = border
        ws.row_dimensions[row].height = 18

    # Legend
    leg_row = len(data["employees"]) + 4
    ws.cell(row=leg_row, column=1, value="Обозначения:").font = bold_dark
    for col, (code, label, fill) in enumerate([
        ("Р","Рабочий день", fill_work),
        ("В","Выходной", fill_weekend),
        ("О","Отпуск", fill_vac),
        ("Б","Больничный", fill_sick),
    ], start=2):
        c = ws.cell(row=leg_row, column=col, value=f"{code} — {label}")
        c.fill = fill
        c.border = border

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"schedule_{store_id}_{month}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


# ─── Продажи ─────────────────────────────────────────────────────────────────
DIVISION_NAME = "Москва 2 КК"

def _safe_float(val, default=0.0):
    try:
        v = float(val)
        return v if v == v else default  # NaN check
    except Exception:
        return default

def _parse_sales_xlsx(content: bytes, filename: str) -> dict:
    """Парсит файл продаж KIDS_МЕСЯЦ.xlsx. Автоопределение листа и строк данных."""
    import re as _re

    # ── 1. Читаем лист ───────────────────────────────────────────────────────
    xl = pd.ExcelFile(io.BytesIO(content))
    sheet_names = xl.sheet_names
    # Ищем лист с "КИД" или "KID" в названии, иначе берём первый
    sheet = None
    for s in sheet_names:
        if "КИД" in s.upper() or "KID" in s.upper():
            sheet = s
            break
    if sheet is None:
        sheet = sheet_names[0]
    try:
        df = pd.read_excel(io.BytesIO(content), sheet_name=sheet, header=None)
    except Exception as e:
        return {"ok": False, "error": f"Ошибка чтения листа '{sheet}': {e}. Листы в файле: {sheet_names}"}

    # ── 2. Дата отчёта ───────────────────────────────────────────────────────
    report_date = None
    for ri in range(min(5, len(df))):
        for ci in range(min(5, len(df.columns))):
            try:
                cell = str(df.iloc[ri, ci])
                m = _re.search(r'(\d{2}\.\d{2}\.\d{4})', cell)
                if m:
                    report_date = m.group(1)
                    break
            except Exception:
                pass
        if report_date:
            break
    if not report_date:
        report_date = datetime.now().strftime("%d.%m.%Y")

    # ── 3. Находим строку-заголовок (ищем по ключевым словам) ───────────────
    header_row = 1   # дефолт
    data_start = 2
    col_map = {}     # имя → индекс

    # Ключевые слова для поиска колонок
    COL_KEYS = {
        "to_fact":     ["то факт", "to_fact", "продажи факт", "факт продаж", "оборот факт",
                        "товарооборот факт", "тов.об.факт", "розн.оборот", "выручка факт",
                        "факт тo", "факт то", "объем продаж"],
        "plan_pct":    ["% план", "% выполн", "план %", "выполнение план", "% от план",
                        "исп.план", "% исп", "вып.план", "выполн.%", "исполн.план",
                        "% вып", "% выполнения", "план выполн"],
        "lfl":         ["lfl", "лфл", "like for like", "l-f-l", "лайк фор лайк"],
        "margin":      ["маржа", "margin", "рент", "мд%", "мд %"],
        "avg_ticket":  ["средн", "avg", "средний чек", "ср.чек", "ср чек"],
        "conversion":  ["конверс", "conversion", "конв."],
        "upt":         ["upt", "ед/чек", "единиц", "ед. на чек"],
        "traffic":     ["трафик", "traffic", "посетит", "кол-во чеков", "количество чеков"],
        "traffic_lfl": ["трафик lfl", "traffic lfl", "трафик лфл"],
        "toys_to":     ["игрушк", "toy", "игр.то", "игр то"],
        "toys_lfl":    ["игрушк lfl", "toy lfl", "игрушк лфл", "игр.lfl", "игр.лфл"],
        "clothes_to":  ["одежд", "cloth", "одеж.то", "одеж то"],
        "clothes_lfl": ["одежд lfl", "одежд лфл", "одеж.lfl", "одеж.лфл"],
        "shoes_to":    ["обувь", "shoe", "обув.то", "обув то"],
        "shoes_lfl":   ["обувь lfl", "обувь лфл", "обув.lfl", "обув.лфл"],
        "sport_to":    ["спорт", "sport", "спорт.то", "спорт то"],
        "sport_lfl":   ["спорт lfl", "спорт лфл", "спорт.lfl", "спорт.лфл"],
    }

    for ri in range(min(10, len(df))):
        row_vals = [str(v).lower().strip() for v in df.iloc[ri]]
        matches = 0
        tmp_map = {}
        for field, keys in COL_KEYS.items():
            for ci, cell in enumerate(row_vals):
                if any(k in cell for k in keys) and field not in tmp_map:
                    tmp_map[field] = ci
                    matches += 1
                    break
        if matches >= 4:
            header_row = ri
            data_start = ri + 1
            col_map = tmp_map
            break

    # Fallback-позиции для полей, не найденных автоопределением
    FALLBACK_COLS = {
        "to_fact": 4, "plan_pct": 5, "lfl": 6, "margin": 8,
        "avg_ticket": 10, "conversion": 12, "upt": 14,
        "traffic": 16, "traffic_lfl": 17,
        "toys_to": 18, "toys_lfl": 19,
        "clothes_to": 21, "clothes_lfl": 22,
        "shoes_to": 24, "shoes_lfl": 25,
        "sport_to": 27, "sport_lfl": 28,
    }
    for field, idx in FALLBACK_COLS.items():
        if field not in col_map:
            col_map[field] = idx

    # Сохраняем debug-информацию о последнем парсинге
    global _last_sales_parse_debug
    _last_sales_parse_debug = {
        "sheet": sheet,
        "header_row": header_row,
        "col_map": col_map,
        "raw_headers": [str(v) for v in df.iloc[header_row]] if header_row < len(df) else [],
    }

    def gcol(row, field, default=0.0):
        idx = col_map.get(field)
        if idx is None or idx >= len(row):
            return default
        return _safe_float(row[idx], default)

    data_df = df.iloc[data_start:].reset_index(drop=True)

    # ── 4. Записываем в БД ───────────────────────────────────────────────────
    conn = get_db()
    conn.execute("DELETE FROM sales_data WHERE report_date=%s", (report_date,))
    conn.commit()

    rows_inserted = 0
    division_row_saved = False

    for _, row in data_df.iterrows():
        try:
            subdivision = str(row.iloc[1] if len(row) > 1 else "").strip()
            store_raw   = str(row.iloc[2] if len(row) > 2 else "").strip()
        except Exception:
            continue

        if not subdivision or subdivision == "nan":
            continue

        is_division = (DIVISION_NAME in subdivision) and (
            not store_raw or store_raw == "nan" or store_raw == subdivision
        )

        store_id = None
        if store_raw and store_raw != "nan":
            try:
                store_id = int(float(store_raw))
                if store_id not in STORE_DATA:
                    store_id = None
            except Exception:
                store_id = None

        if not is_division and store_id is None:
            continue

        plan_pct = gcol(row, "plan_pct")
        # Если значение в долях (0.9542 → 95.42%)
        if 0 < plan_pct <= 5:
            plan_pct = round(plan_pct * 100, 2)
        # Если > 200, скорее всего промиле или ошибка — обнуляем
        if plan_pct > 200:
            plan_pct = 0.0

        def norm_lfl(v):
            if -5 < v < 5 and v != 0:
                return round(v * 100, 2)
            return round(v, 2)

        lfl = norm_lfl(gcol(row, "lfl"))
        toys_lfl    = norm_lfl(gcol(row, "toys_lfl"))
        clothes_lfl = norm_lfl(gcol(row, "clothes_lfl"))
        shoes_lfl   = norm_lfl(gcol(row, "shoes_lfl"))
        sport_lfl   = norm_lfl(gcol(row, "sport_lfl"))
        traffic_lfl = norm_lfl(gcol(row, "traffic_lfl"))

        traffic_raw = gcol(row, "traffic")
        traffic = int(traffic_raw) if traffic_raw else 0

        conn.execute("""
            INSERT INTO sales_data
              (report_date, store_id, subdivision, to_fact, plan_pct, lfl, margin,
               avg_ticket, conversion, upt, traffic, traffic_lfl,
               toys_to, toys_lfl, clothes_to, clothes_lfl,
               shoes_to, shoes_lfl, sport_to, sport_lfl)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
        """, (report_date,
              None if is_division else store_id,
              DIVISION_NAME if is_division else subdivision,
              gcol(row, "to_fact"), plan_pct, lfl, gcol(row, "margin"),
              gcol(row, "avg_ticket"), gcol(row, "conversion"), gcol(row, "upt"),
              traffic, traffic_lfl,
              gcol(row, "toys_to"), toys_lfl,
              gcol(row, "clothes_to"), clothes_lfl,
              gcol(row, "shoes_to"), shoes_lfl,
              gcol(row, "sport_to"), sport_lfl))
        rows_inserted += 1
        if is_division:
            division_row_saved = True

    conn.execute("INSERT INTO sales_upload_log (filename, report_date, rows_count) VALUES (%s,%s,%s)",
                 (filename, report_date, rows_inserted))
    conn.commit()
    conn.close()
    return {"ok": True, "report_date": report_date, "rows": rows_inserted,
            "division_saved": division_row_saved, "sheet": sheet, "col_map": col_map}

@app.post("/api/sales/upload")
async def upload_sales(file: UploadFile = File(...), user=Depends(require_admin)):
    content = await file.read()
    result = _parse_sales_xlsx(content, file.filename)
    if not result.get("ok"):
        raise HTTPException(status_code=400, detail=result["error"])
    return result

@app.get("/api/sales/store/{store_id}")
def get_sales_store(store_id: int, user=Depends(require_auth)):
    # Директор может смотреть только свой магазин
    if user["role"] != "admin" and user.get("store_id") != store_id:
        raise HTTPException(status_code=403, detail="Нет доступа")
    conn = get_db()
    # Последняя дата
    last = conn.execute(
        "SELECT report_date FROM sales_data ORDER BY uploaded_at DESC LIMIT 1"
    ).fetchone()
    if not last:
        conn.close()
        return {"store": None, "division": None, "report_date": None}
    rdate = last["report_date"]
    store = conn.execute(
        "SELECT * FROM sales_data WHERE report_date=%s AND store_id=%s",
        (rdate, store_id)
    ).fetchone()
    division = conn.execute(
        "SELECT * FROM sales_data WHERE report_date=%s AND store_id IS NULL",
        (rdate,)
    ).fetchone()
    conn.close()
    return {
        "report_date": rdate,
        "store": dict(store) if store else None,
        "division": dict(division) if division else None
    }

@app.get("/api/sales/all")
def get_sales_all(user=Depends(require_admin)):
    conn = get_db()
    last = conn.execute(
        "SELECT report_date FROM sales_data ORDER BY uploaded_at DESC LIMIT 1"
    ).fetchone()
    if not last:
        conn.close()
        return {"stores": [], "division": None, "report_date": None}
    rdate = last["report_date"]
    stores = conn.execute(
        "SELECT * FROM sales_data WHERE report_date=%s AND store_id IS NOT NULL ORDER BY to_fact DESC",
        (rdate,)
    ).fetchall()
    division = conn.execute(
        "SELECT * FROM sales_data WHERE report_date=%s AND store_id IS NULL",
        (rdate,)
    ).fetchone()
    conn.close()
    return {
        "report_date": rdate,
        "stores": [dict(r) for r in stores],
        "division": dict(division) if division else None
    }

@app.get("/api/sales/log")
def sales_log(user=Depends(require_admin)):
    conn = get_db()
    rows = conn.execute(
        "SELECT * FROM sales_upload_log ORDER BY uploaded_at DESC LIMIT 10"
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]

@app.get("/api/sales/debug-headers")
def sales_debug_headers():
    """Показывает заголовки Excel и определённые колонки из последнего парсинга."""
    return _last_sales_parse_debug

# ─── Остатки (Stock) ──────────────────────────────────────────────────────────

def _parse_stock_xlsx(content: bytes, filename: str) -> dict:
    import re as _re
    xl = pd.ExcelFile(io.BytesIO(content))
    sheet_names = xl.sheet_names
    sheet = None
    for s in sheet_names:
        if any(k in s.upper() for k in ["КИД", "KID", "ОСТАТ", "STOCK", "ЗАПАС"]):
            sheet = s
            break
    if sheet is None:
        sheet = sheet_names[0]
    try:
        df = pd.read_excel(io.BytesIO(content), sheet_name=sheet, header=None)
    except Exception as e:
        return {"ok": False, "error": f"Ошибка чтения листа '{sheet}': {e}"}

    report_date = None
    for ri in range(min(5, len(df))):
        for ci in range(min(5, len(df.columns))):
            try:
                cell = str(df.iloc[ri, ci])
                m = _re.search(r'(\d{2}\.\d{2}\.\d{4})', cell)
                if m:
                    report_date = m.group(1)
                    break
            except Exception:
                pass
        if report_date:
            break
    if not report_date:
        report_date = datetime.now().strftime("%d.%m.%Y")

    STOCK_KEYS = {
        "total_stock":   ["итого", "всего", "total", "общий", "остат.итог", "остатки итого", "сумма остат"],
        "total_lfl":     ["итого lfl", "всего lfl", "total lfl", "итого лфл"],
        "clothes_stock": ["одежд", "cloth"],
        "clothes_lfl":   ["одежд lfl", "одежд лфл", "одеж.lfl"],
        "toys_stock":    ["игрушк", "toy", "игр."],
        "toys_lfl":      ["игрушк lfl", "toy lfl", "игр.lfl", "игр.лфл"],
        "shoes_stock":   ["обувь", "shoe", "обув."],
        "shoes_lfl":     ["обувь lfl", "обувь лфл", "обув.lfl"],
        "sport_stock":   ["спорт", "sport"],
        "sport_lfl":     ["спорт lfl", "спорт лфл", "спорт.lfl"],
    }

    col_map = {}
    data_start = 1
    for ri in range(min(10, len(df))):
        row_vals = [str(v).lower().strip() for v in df.iloc[ri]]
        matches = 0
        tmp_map = {}
        for field, keys in STOCK_KEYS.items():
            for ci, cell in enumerate(row_vals):
                if any(k in cell for k in keys) and field not in tmp_map:
                    tmp_map[field] = ci
                    matches += 1
                    break
        if matches >= 3:
            data_start = ri + 1
            col_map = tmp_map
            break

    if not col_map:
        col_map = {
            "total_stock": 3, "total_lfl": 4,
            "clothes_stock": 5, "clothes_lfl": 6,
            "toys_stock": 7, "toys_lfl": 8,
            "shoes_stock": 9, "shoes_lfl": 10,
            "sport_stock": 11, "sport_lfl": 12,
        }

    def gcol(row, field, default=0.0):
        idx = col_map.get(field)
        if idx is None or idx >= len(row):
            return default
        return _safe_float(row[idx], default)

    def norm_lfl(v):
        if -5 < v < 5 and v != 0:
            return round(v * 100, 2)
        return round(v, 2)

    data_df = df.iloc[data_start:].reset_index(drop=True)
    conn = get_db()
    conn.execute("DELETE FROM stock_data WHERE report_date=%s", (report_date,))
    conn.commit()

    rows_inserted = 0
    for _, row in data_df.iterrows():
        try:
            subdivision = str(row.iloc[1] if len(row) > 1 else "").strip()
            store_raw   = str(row.iloc[2] if len(row) > 2 else "").strip()
        except Exception:
            continue
        if not subdivision or subdivision == "nan":
            continue

        is_division = (DIVISION_NAME in subdivision) and (
            not store_raw or store_raw == "nan" or store_raw == subdivision
        )
        store_id = None
        if store_raw and store_raw != "nan":
            try:
                store_id = int(float(store_raw))
                if store_id not in STORE_DATA:
                    store_id = None
            except Exception:
                store_id = None

        if not is_division and store_id is None:
            continue

        conn.execute("""
            INSERT INTO stock_data
              (report_date, store_id, subdivision,
               total_stock, total_lfl,
               clothes_stock, clothes_lfl,
               toys_stock, toys_lfl,
               shoes_stock, shoes_lfl,
               sport_stock, sport_lfl)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
        """, (report_date,
              None if is_division else store_id,
              DIVISION_NAME if is_division else subdivision,
              gcol(row, "total_stock"),   norm_lfl(gcol(row, "total_lfl")),
              gcol(row, "clothes_stock"), norm_lfl(gcol(row, "clothes_lfl")),
              gcol(row, "toys_stock"),    norm_lfl(gcol(row, "toys_lfl")),
              gcol(row, "shoes_stock"),   norm_lfl(gcol(row, "shoes_lfl")),
              gcol(row, "sport_stock"),   norm_lfl(gcol(row, "sport_lfl"))))
        rows_inserted += 1

    conn.execute("INSERT INTO stock_upload_log (filename, report_date, rows_count) VALUES (%s,%s,%s)",
                 (filename, report_date, rows_inserted))
    conn.commit()
    conn.close()
    return {"ok": True, "report_date": report_date, "rows": rows_inserted, "sheet": sheet}

@app.post("/api/stock/upload")
async def upload_stock(file: UploadFile = File(...), user=Depends(require_admin)):
    content = await file.read()
    result = _parse_stock_xlsx(content, file.filename)
    if not result.get("ok"):
        raise HTTPException(status_code=400, detail=result["error"])
    return result

@app.get("/api/stock/store/{store_id}")
def get_stock_store(store_id: int, user=Depends(require_auth)):
    if user["role"] != "admin" and user.get("store_id") != store_id:
        raise HTTPException(status_code=403, detail="Нет доступа")
    conn = get_db()
    last = conn.execute(
        "SELECT report_date FROM stock_data ORDER BY uploaded_at DESC LIMIT 1"
    ).fetchone()
    if not last:
        conn.close()
        return {"store": None, "division": None, "report_date": None}
    rdate = last["report_date"]
    store = conn.execute(
        "SELECT * FROM stock_data WHERE report_date=%s AND store_id=%s",
        (rdate, store_id)
    ).fetchone()
    division = conn.execute(
        "SELECT * FROM stock_data WHERE report_date=%s AND store_id IS NULL",
        (rdate,)
    ).fetchone()
    conn.close()
    return {
        "report_date": rdate,
        "store": dict(store) if store else None,
        "division": dict(division) if division else None
    }

@app.get("/api/stock/all")
def get_stock_all(user=Depends(require_admin)):
    conn = get_db()
    last = conn.execute(
        "SELECT report_date FROM stock_data ORDER BY uploaded_at DESC LIMIT 1"
    ).fetchone()
    if not last:
        conn.close()
        return {"stores": [], "division": None, "report_date": None}
    rdate = last["report_date"]
    stores = conn.execute(
        "SELECT * FROM stock_data WHERE report_date=%s AND store_id IS NOT NULL ORDER BY total_stock DESC",
        (rdate,)
    ).fetchall()
    division = conn.execute(
        "SELECT * FROM stock_data WHERE report_date=%s AND store_id IS NULL",
        (rdate,)
    ).fetchone()
    conn.close()
    return {
        "report_date": rdate,
        "stores": [dict(r) for r in stores],
        "division": dict(division) if division else None
    }

@app.get("/api/stock/log")
def stock_log(user=Depends(require_admin)):
    conn = get_db()
    rows = conn.execute(
        "SELECT * FROM stock_upload_log ORDER BY uploaded_at DESC LIMIT 10"
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]

# ─── Email авто-загрузка мотивации ───────────────────────────────────────────
def _decode_filename(raw):
    parts = decode_email_header(raw)
    result = ""
    for part, enc in parts:
        if isinstance(part, bytes):
            result += part.decode(enc or "utf-8", errors="replace")
        else:
            result += part
    return result

def fetch_motivation_from_email() -> dict:
    """Подключается к почте, находит последнее письмо с файлом мотивации и загружает его."""
    if not EMAIL_USER or not EMAIL_PASSWORD:
        return {"ok": False, "error": "EMAIL_USER / EMAIL_PASSWORD не настроены в Railway Variables"}
    try:
        mail = imaplib.IMAP4_SSL(EMAIL_HOST, EMAIL_PORT)
        mail.login(EMAIL_USER, EMAIL_PASSWORD)
        mail.select("INBOX")

        # Ищем по ключевому слову темы (работает даже с Fwd: Мотивация ...)
        keyword = EMAIL_SUBJECT.split()[0]  # например 'Мотивация'
        _, msgs = mail.search(None, f'SUBJECT "{keyword}"')
        ids = msgs[0].split() if msgs[0] else []

        # Если не нашли по SUBJECT IMAP — ищем среди ALL и фильтруем вручную
        if not ids:
            _, msgs = mail.search(None, 'ALL')
            ids = msgs[0].split() if msgs[0] else []

        # Берём последнее письмо, тема которого содержит ключевое слово
        target_id = None
        for eid in reversed(ids):
            _, hdr = mail.fetch(eid, '(BODY[HEADER.FIELDS (SUBJECT)])')
            subj_raw = hdr[0][1].decode('utf-8', errors='replace') if hdr[0] else ''
            if keyword.lower() in subj_raw.lower():
                target_id = eid
                break

        if not target_id:
            mail.logout()
            return {'ok': False, 'error': f'Письмо с темой {EMAIL_SUBJECT!r} не найдено в {EMAIL_USER}'}

        _, msg_data = mail.fetch(target_id, "(RFC822)")
        mail.logout()
        msg = email_lib.message_from_bytes(msg_data[0][1])

        # Ищем xlsx-вложение
        attachment_data = None
        filename = "motivation.xlsx"
        for part in msg.walk():
            if part.get_content_disposition() in ("attachment", "inline"):
                raw_name = part.get_filename()
                if raw_name:
                    fname = _decode_filename(raw_name)
                    if fname.lower().endswith((".xlsx", ".xls")):
                        attachment_data = part.get_payload(decode=True)
                        filename = fname
                        break

        if not attachment_data:
            return {"ok": False, "error": "xlsx-вложение не найдено в письме"}

        # Обрабатываем файл как загрузку мотивации
        result = _process_motivation_bytes(attachment_data, filename)
        result["source"] = "email"
        return result

    except Exception as e:
        return {"ok": False, "error": str(e)}

def _process_motivation_bytes(content: bytes, filename: str) -> dict:
    """Обрабатывает xlsx-данные мотивации (общий код для upload и email)."""
    try:
        df = pd.read_excel(io.BytesIO(content), header=None)
    except Exception as e:
        return {"ok": False, "error": f"Ошибка чтения файла: {e}"}

    header_row = None
    for i, row in df.iterrows():
        if str(row[0]).strip().startswith("Дата"):
            header_row = i
            break
    if header_row is None:
        return {"ok": False, "error": "Не найден заголовок 'Дата отчета'"}

    data_df = df.iloc[header_row + 1:].reset_index(drop=True)
    conn = get_db()
    rows_inserted = 0
    report_date = None

    try:
        # Получим все store_ids из файла для удаления дублей (без дубля первого DELETE)
        store_ids_in_file = set()
        for _, row in data_df.iterrows():
            try:
                sid = int(float(str(row[2]))) if str(row[2]) not in ["nan", ""] else None
                if sid:
                    store_ids_in_file.add(sid)
            except Exception:
                continue

        if store_ids_in_file:
            placeholders = ",".join(["%s"] * len(store_ids_in_file))
            conn.execute(f"DELETE FROM motivation_data WHERE store_id IN ({placeholders})", tuple(store_ids_in_file))
            conn.commit()
    except Exception as e:
        conn.rollback()
        conn.close()
        return {"ok": False, "error": f"Ошибка очистки старых данных: {e}"}

    for _, row in data_df.iterrows():
        try:
            store_id = int(float(str(row[2]))) if str(row[2]) not in ["nan", ""] else None
            if not store_id:
                continue
            login   = str(row[3] or "").strip()
            role    = str(row[5] or "").strip()
            name    = str(row[6] or "").strip()
            to_fact = float(row[7]) if str(row[7]) not in ["nan", ""] else 0
            pct_to  = float(row[8]) if str(row[8]) not in ["nan", ""] else 0
            income  = float(row[9]) if str(row[9]) not in ["nan", ""] else 0
            fdm_val = float(row[10]) if str(row[10]) not in ["nan", ""] else 0
            raw_date = row[0]
            if hasattr(raw_date, "strftime"):
                rd = raw_date.strftime("%d.%m.%Y")
            else:
                rd = str(raw_date).split(" ")[0].strip()
            if report_date is None:
                report_date = rd
            is_total   = 1 if login == "Total" else 0
            is_bezshk  = 1 if login == "БезШК" else 0
            conn.execute("""
                INSERT INTO motivation_data
                    (store_id, report_date, login, role, name, to_fact, pct_to, income, fdm, is_total, is_bezshk)
                VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
            """, (store_id, rd, login, role, name, to_fact, pct_to, income, fdm_val, is_total, is_bezshk))
            rows_inserted += 1
        except Exception:
            conn.rollback()  # сбрасываем ошибку транзакции PostgreSQL перед следующей строкой
            continue

    try:
        conn.execute("INSERT INTO upload_log (filename, report_date, rows_count) VALUES (%s,%s,%s)",
                     (filename, report_date, rows_inserted))
        conn.commit()
    except Exception:
        conn.rollback()
    finally:
        conn.close()
    return {"ok": True, "filename": filename, "rows": rows_inserted, "report_date": report_date}


@app.get("/api/email/fetch-now")
def email_fetch_now(user=Depends(require_admin)):
    """Ручной запуск загрузки файла мотивации из почты."""
    result = fetch_motivation_from_email()
    if not result.get("ok"):
        raise HTTPException(status_code=500, detail=result.get("error"))
    return result

@app.get("/api/email/status")
def email_status(user=Depends(require_admin)):
    """Статус настройки автозагрузки."""
    return {
        "configured": bool(EMAIL_USER and EMAIL_PASSWORD),
        "host": EMAIL_HOST,
        "user": EMAIL_USER or "не задан",
        "sender_filter": EMAIL_SENDER,
        "subject_filter": EMAIL_SUBJECT,
        "schedule": f"ежедневно в {EMAIL_HOUR:02d}:{EMAIL_MINUTE:02d}",
        "last_run": _last_email_fetch_result
    }

# ─── Планировщик (встроенный threading, без внешних зависимостей) ─────────────
_last_email_fetch_result: dict = {"status": "never", "time": None, "error": None, "rows": None}
_last_sales_parse_debug: dict = {}

def _email_scheduler_loop():
    """Фоновый поток: проверяет время и запускает загрузку раз в день."""
    global _last_email_fetch_result
    last_run_date = None
    while True:
        try:
            now = datetime.utcnow() + timedelta(hours=3)  # MSK
            if now.hour == EMAIL_HOUR and now.minute == EMAIL_MINUTE and now.date() != last_run_date:
                last_run_date = now.date()
                result = fetch_motivation_from_email()
                _last_email_fetch_result = {
                    "status": "ok" if result.get("ok") else "error",
                    "time": now.strftime("%d.%m.%Y %H:%M"),
                    "error": result.get("error"),
                    "rows": result.get("rows"),
                    "filename": result.get("filename"),
                }
                print(f"[Email Auto-Fetch] {now.isoformat()} → {result}")
        except Exception as e:
            _last_email_fetch_result = {
                "status": "error", "time": None, "error": str(e), "rows": None, "filename": None
            }
            print(f"[Email Auto-Fetch] Ошибка планировщика: {e}")
        time.sleep(30)  # Проверяем каждые 30 секунд

_scheduler_thread = threading.Thread(target=_email_scheduler_loop, daemon=True)
_scheduler_thread.start()

# ─── Отдаём PWA ──────────────────────────────────────────────────────────────
app.mount("/static", StaticFiles(directory="../frontend/static"), name="static")

@app.get("/{full_path:path}")
def serve_spa(full_path: str):
    return FileResponse("../frontend/index.html")
